"""Tests pytest de la logique pure v5 (pas de Streamlit)."""
import datetime as dt
import sys, os

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import db_gen
import integrity
import rapport

TODAY = dt.date(2026, 10, 8)


def _base_trois_mois():
    data = db_gen.base_vide()
    for a, m in ((2026, 10), (2026, 8), (2026, 11)):
        data, _, _ = db_gen.ajouter_mois(data, a, m)
    return data


def test_generation_et_renumerotation():
    data = _base_trois_mois()
    assert len(data) == 12
    assert [r["n"] for r in data] == list(range(1, 13))
    assert [r["date_echeance"] for r in data] == sorted(r["date_echeance"] for r in data)


def test_anti_doublon():
    data, ok, _ = db_gen.ajouter_mois(db_gen.base_vide(), 2026, 10)
    data, ok2, _ = db_gen.ajouter_mois(data, 2026, 10)
    assert ok and not ok2


def test_fevrier_28():
    data, _, _ = db_gen.ajouter_mois(db_gen.base_vide(), 2027, 2)
    s3 = [r for r in data if r["type"] == db_gen.T_S3][0]
    assert s3["date_echeance"] == "2027-02-28"


def test_regularite_au_jour():
    # payé le 28 un dû le 10 (même mois) → retard
    assert integrity.regularite("2026-10-28", "2026-10-10") == "retard"
    # payé le 05 pour échéance le 10 (même mois) → à temps
    assert integrity.regularite("2026-10-05", "2026-10-10") == "à temps"
    # payé en septembre pour échéance de novembre → avance
    assert integrity.regularite("2026-09-20", "2026-11-10") == "avance"
    # payé le jour même → à temps
    assert integrity.regularite("2026-10-10", "2026-10-10") == "à temps"


def test_en_retard_badge():
    data = _base_trois_mois()
    aout = [r for r in data if r["mois"] == "Août 2026" and r["type"] == db_gen.T_S1][0]
    assert integrity.est_en_retard(aout, TODAY) is True
    aout["statut"] = "Payé – retard"
    aout["date_paiement"] = "2026-09-01"
    assert integrity.est_en_retard(aout, TODAY) is False


def test_reconcile_repare_montant_mois_et_statut():
    sale = [
        {"date_echeance": "2026-08-10", "mois": "FAUX", "type": db_gen.T_CN,
         "montant": 999, "statut": "Payé – à temps", "rappel_j1": "X",
         "observations": "", "date_paiement": "2026-09-10"},   # payé après échéance → retard
        {"foo": "bar"},                                          # invalide
    ]
    clean, rep = integrity.reconcile(sale, TODAY)
    assert rep["rejetes"] == 1
    r = clean[0]
    assert r["montant"] == 1500               # corrigé au barème
    assert r["mois"] == "Août 2026"           # corrigé depuis la date
    assert r["rappel_j1"] == "2026-08-09"     # corrigé
    assert r["statut"] == "Payé – retard"     # recalculé au jour


def test_reconcile_coherence_non_paye():
    rec = db_gen.generer_mois(2026, 10)[0]
    rec["date_paiement"] = "2026-10-10"       # incohérent : non payé avec date
    clean, rep = integrity.reconcile([rec], TODAY)
    assert clean[0]["date_paiement"] is None


def test_export_excel_8_colonnes_et_total():
    data = _base_trois_mois()
    data[0]["statut"] = "Payé – à temps"
    data[0]["date_paiement"] = data[0]["date_echeance"]
    xb = rapport.build_recap_xlsx(data, only_paid=True)
    from openpyxl import load_workbook
    from io import BytesIO
    ws = load_workbook(BytesIO(xb)).active
    assert [c.value for c in ws[1]] == rapport.COLONNES
    assert ws.cell(row=ws.max_row, column=5).value == data[0]["montant"]


def test_ics_rappels_non_payes():
    data = _base_trois_mois()
    ics = rapport.build_ics(data, only_unpaid=True).decode("utf-8")
    assert ics.count("BEGIN:VEVENT") == 12
    assert ics.count("TRIGGER:-PT15H") == 12
    assert ics.startswith("BEGIN:VCALENDAR")
